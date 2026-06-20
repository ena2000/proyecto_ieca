"""
Inserta Metodología de investigación en CAPITULO 3 - ENTENDERLO.docx
(antes de Metodología de desarrollo y verificación).
"""
from __future__ import annotations

import importlib.util
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

SCRIPT_DIR = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location(
    "metodologia_investigacion_contenido",
    SCRIPT_DIR / "metodologia-investigacion-contenido.py",
)
_mod = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_mod)

DOC_PATH = Path(r"c:\Users\ena\Desktop\TITULACION\CAPITULO 3 - ENTENDERLO.docx")
OUT_FALLBACK = DOC_PATH.with_name(DOC_PATH.stem + "-metodologia-inv.docx")
HEADING = _mod.HEADING
ANCHOR = "Metodología de desarrollo y verificación"
STOP_BEFORE = "Casos de uso integrados (Anexo 18)"


def set_run_font(run, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(text.replace("**", "")), bold=bold, size=size)


def find_heading(doc: Document, text: str) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    return None


def paragraph_text(element) -> str:
    texts = [t.text for t in element.iter(qn("w:t")) if t.text]
    return "".join(texts).strip()


def find_paragraph_element(doc: Document, text: str) -> CT_P | None:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p._element
    return None


def remove_before_anchor(doc: Document, anchor_text: str, stop_before: str) -> None:
    """Elimina párrafos y tablas entre stop_before y anchor_text (sin borrar stop_before)."""
    anchor = find_paragraph_element(doc, anchor_text)
    if anchor is None:
        raise ValueError(f"No se encontró ancla: {anchor_text!r}")

    while True:
        prev = anchor.getprevious()
        if prev is None:
            break
        if prev.tag == qn("w:p") and paragraph_text(prev) == stop_before:
            break
        prev.getparent().remove(prev)


def make_heading(doc: Document, text: str) -> Paragraph:
    p = doc.add_heading(text, level=3)
    for run in p.runs:
        set_run_font(run, bold=True, size=12)
    return p


def make_para(doc: Document, text: str, bold: bool = False) -> Paragraph:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text.replace("**", "")), bold=bold, size=11)
    return p


def make_table(doc: Document, rows: list[list[str]]) -> Table:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val, bold=(r == 0), size=9)
    return table


def insert_before(anchor, element) -> None:
    anchor.addprevious(element)


def build_blocks(doc: Document) -> list:
    blocks = []
    blocks.append(("h", _mod.METODOLOGIAS_PROYECTO_HEADING))
    for p in _mod.METODOLOGIAS_PROYECTO:
        blocks.append(("p", p))
    blocks.append(("h", HEADING))
    for p in _mod.PARAGRAPHS:
        blocks.append(("p", p))
    blocks.append(("h", _mod.POBLACION_TITLE))
    blocks.append(("p", _mod.POBLACION))
    blocks.append(("pb", _mod.TABLA_POBLACION_CAPTION))
    blocks.append(("t", _mod.TABLA_POBLACION))
    blocks.append(("p", _mod.NOTA_POBLACION))
    blocks.append(("p", _mod.MUESTRA_INTRO))
    blocks.append(("p", _mod.MUESTRA_PARAMS))
    for p in _mod.MUESTRA_FORMULA:
        blocks.append(("p", p))
    for p in _mod.MUESTRA_FRACCION:
        blocks.append(("p", p))
    blocks.append(("p", _mod.MUESTRA_APLICADA))
    blocks.append(("pb", _mod.TABLA_MUESTRA_CAPTION))
    blocks.append(("t", _mod.TABLA_MUESTRA))
    blocks.append(("p", _mod.NOTA_MUESTRA))
    blocks.append(("h", _mod.PROCESAMIENTO_TITLE))
    for p in _mod.PROCESAMIENTO:
        blocks.append(("p", p))
    blocks.append(("h", _mod.TECNICAS_TITLE))
    blocks.append(("t", _mod.TECNICAS))
    blocks.append(("pb", _mod.PREGUNTA_EJEMPLO))
    blocks.append(("pb", _mod.TABLA_PREGUNTA_CAPTION))
    blocks.append(("t", _mod.TABLA_PREGUNTA))
    blocks.append(("p", _mod.FIGURA_EJEMPLO))
    blocks.append(("p", _mod.ANALISIS_PREGUNTA))
    blocks.append(("pb", _mod.TABLA_PRUEBAS_CAPTION))
    blocks.append(("t", _mod.TABLA_PRUEBAS))
    blocks.append(("p", _mod.ANALISIS_PRUEBAS))
    return blocks


def insert_section(doc: Document) -> None:
    idx = find_heading(doc, ANCHOR)
    if idx is None:
        raise ValueError(f"No se encontró: {ANCHOR!r}")
    anchor = doc.paragraphs[idx]._element

    for kind, data in build_blocks(doc):
        if kind == "h":
            el = make_heading(doc, data)._element
        elif kind == "pb":
            el = make_para(doc, data, bold=True)._element
        elif kind == "p":
            el = make_para(doc, data)._element
        elif kind == "t":
            el = make_table(doc, data)._tbl
        else:
            continue
        insert_before(anchor, el)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup = DOC_PATH.with_suffix(
        f".backup-met-inv-{datetime.now():%Y%m%d-%H%M%S}.docx"
    )
    shutil.copy2(DOC_PATH, backup)
    print(f"Respaldo: {backup}")

    doc = Document(str(DOC_PATH))
    remove_before_anchor(doc, ANCHOR, STOP_BEFORE)
    insert_section(doc)

    try:
        doc.save(str(DOC_PATH))
        print(f"Guardado: {DOC_PATH}")
    except PermissionError:
        doc.save(str(OUT_FALLBACK))
        print(f"Archivo abierto. Guardado en: {OUT_FALLBACK}")


if __name__ == "__main__":
    main()
