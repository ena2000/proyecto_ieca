"""
Agrega Entregables del proyecto y Propuesta al final de CAPITULO 3 - ENTENDERLO.docx.
"""
from __future__ import annotations

import importlib.util
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

SCRIPT_DIR = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location(
    "entregables_propuesta_contenido",
    SCRIPT_DIR / "entregables-propuesta-contenido.py",
)
_mod = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_mod)

DOC_PATH = Path(r"c:\Users\ena\Desktop\TITULACION\CAPITULO 3 - ENTENDERLO.docx")
OUT_FALLBACK = DOC_PATH.with_name(DOC_PATH.stem + "-entregables-propuesta.docx")

HEADINGS = (_mod.HEADING_ENTREGABLES, _mod.HEADING_PROPUESTA)


def set_run_font(run, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(text), bold=bold, size=size)


def add_heading(doc: Document, text: str, level: int = 3) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, size=12)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), bold=bold, size=11)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val, bold=(r == 0), size=9)
    doc.add_paragraph()


def find_heading_index(doc: Document, text: str) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    return None


def remove_sections_from(doc: Document, start_text: str) -> None:
    start = find_heading_index(doc, start_text)
    if start is None:
        return
    for i in range(len(doc.paragraphs) - 1, start - 1, -1):
        doc.paragraphs[i]._element.getparent().remove(doc.paragraphs[i]._element)
    while len(doc.tables) > 12:
        doc.tables[-1]._tbl.getparent().remove(doc.tables[-1]._tbl)


def append_sections(doc: Document) -> None:
    add_heading(doc, _mod.HEADING_ENTREGABLES, level=3)
    add_para(doc, _mod.ENTREGABLES_INTRO)
    add_para(doc, _mod.TABLA_EDT_CAPTION, bold=True)
    add_table(doc, _mod.TABLA_EDT)

    add_heading(doc, _mod.HEADING_PROPUESTA, level=3)
    for text in _mod.PROPUESTA_PARAGRAPHS:
        add_para(doc, text)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup = DOC_PATH.with_suffix(
        f".backup-entregables-{datetime.now():%Y%m%d-%H%M%S}.docx"
    )
    shutil.copy2(DOC_PATH, backup)
    print(f"Respaldo: {backup}")

    doc = Document(str(DOC_PATH))
    for heading in HEADINGS:
        remove_sections_from(doc, heading)
    append_sections(doc)

    try:
        doc.save(str(DOC_PATH))
        print(f"Guardado: {DOC_PATH}")
    except PermissionError:
        doc.save(str(OUT_FALLBACK))
        print(f"Archivo abierto en Word. Guardado en: {OUT_FALLBACK}")


if __name__ == "__main__":
    main()
