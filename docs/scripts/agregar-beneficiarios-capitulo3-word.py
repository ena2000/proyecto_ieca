"""
Actualiza la sección Beneficiarios (prosa narrativa) en CAPITULO 3 - ENTENDERLO.docx.
"""
from __future__ import annotations

import importlib.util
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

SCRIPT_DIR = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location(
    "beneficiarios_contenido",
    SCRIPT_DIR / "beneficiarios-contenido.py",
)
_mod = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_mod)

HEADING = _mod.HEADING
PARAGRAPHS = _mod.PARAGRAPHS

DOC_PATH = Path(r"c:\Users\ena\Desktop\TITULACION\CAPITULO 3 - ENTENDERLO.docx")
OUT_FALLBACK = DOC_PATH.with_name(DOC_PATH.stem + "-beneficiarios.docx")


def set_run_font(run, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading(doc: Document, text: str, level: int = 3) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, size=12)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), size=11)


def remove_existing_beneficiarios(doc: Document) -> None:
    start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == HEADING:
            start = i
            break
    if start is None:
        return
    for i in range(len(doc.paragraphs) - 1, start - 1, -1):
        doc.paragraphs[i]._element.getparent().remove(doc.paragraphs[i]._element)
    while len(doc.tables) > 12:
        doc.tables[-1]._tbl.getparent().remove(doc.tables[-1]._tbl)


def append_beneficiarios(doc: Document) -> None:
    add_heading(doc, HEADING, level=3)
    for text in PARAGRAPHS:
        add_para(doc, text)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup = DOC_PATH.with_suffix(
        f".backup-beneficiarios-{datetime.now():%Y%m%d-%H%M%S}.docx"
    )
    shutil.copy2(DOC_PATH, backup)
    print(f"Respaldo: {backup}")

    doc = Document(str(DOC_PATH))
    remove_existing_beneficiarios(doc)
    append_beneficiarios(doc)

    try:
        doc.save(str(DOC_PATH))
        print(f"Guardado: {DOC_PATH}")
    except PermissionError:
        doc.save(str(OUT_FALLBACK))
        print(f"Archivo abierto en Word. Guardado en: {OUT_FALLBACK}")


if __name__ == "__main__":
    main()
