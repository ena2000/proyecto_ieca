"""
Reemplaza el placeholder de Beneficiarios directos e indirectos en la plantilla de tesis
con el contenido adaptado al proyecto IECA.
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

DOC_PATH = Path(
    r"c:\Users\ena\Desktop\TITULACION\PLANTILLA DE TESIS - MILENA MARISCAL PONCE.docx"
)
OUT_FALLBACK = DOC_PATH.with_name(
    DOC_PATH.stem + "-beneficiarios.docx"
)
STYLE_BODY = "Estilo"
HEADING = "Beneficiarios directos e indirectos del proyecto"
END_HEADING = "Entregables del proyecto"

INTRO = (
    "Los beneficiarios, involucrados o stakeholders de un proyecto son las personas u "
    "organizaciones que obtendrán algún tipo de beneficio de la implementación de este. "
    "Se pueden identificar dos tipos de beneficiarios: directos e indirectos."
)

P_DIRECTOS = (
    "Beneficiarios directos: Los beneficiarios directos son aquéllos que participarán "
    "directamente en el proyecto y, por consiguiente, se beneficiarán de su implementación. "
    "En el Sistema Web de Gestión Financiera IECA, el administrador que aprueba movimientos, "
    "el contable que elabora reportes, los colaboradores que registran ingresos y gastos de "
    "su ministerio, la desarrolladora del proyecto y la Iglesia Evangélica La Alborada como "
    "institución propietaria son beneficiarios directos, pues participan en el desarrollo, "
    "la validación o el uso habitual del producto desplegado en Firebase Hosting y Render."
)

P_INDIRECTOS = (
    "Beneficiarios indirectos: Los beneficiarios indirectos son, con frecuencia, pero no "
    "siempre, las personas que se encuentran dentro de la zona de influencia del proyecto. "
    "Aunque el sistema está diseñado para un número acotado de usuarios operativos, el "
    "beneficio se extiende a la feligresía, el liderazgo pastoral, los ministerios que "
    "reciben asignaciones o reportes, los auditores internos y los futuros administradores "
    "o contables, quienes se benefician de la transparencia, el orden contable y la "
    "documentación que el sistema genera sin operar el panel de forma diaria."
)

P_MARCO = (
    "En esta sección se complementa con la Fase 2 de la Metodología de Marco Lógica "
    "denominada Análisis de involucrados, mediante la identificación y categorización de "
    "actores (usuarios primarios, patrocinador, ejecutor técnico y beneficiarios finales "
    "indirectos) y la matriz de involucrados que relaciona interés, influencia y estrategia "
    "de participación para cada stakeholder del proyecto IECA."
)

TABLA_DIRECTOS = [
    ["Beneficiario", "Rol en el proyecto / sistema", "Beneficio obtenido"],
    [
        "Administrador IECA",
        "Actor principal (CU-06 a CU-15)",
        "Aprobación de movimientos, usuarios, ministerios, cierre, backup y auditoría",
    ],
    [
        "Contable IECA",
        "Actor de consulta (CU-08, CU-09, CU-17)",
        "Reportes, kardex, exportación Excel y alertas; sin aprobar movimientos",
    ],
    [
        "Colaboradores de ministerio",
        "Actores operativos (CU-04, CU-05, CU-18)",
        "Registro digital de ingresos y gastos con comprobantes",
    ],
    [
        "Desarrolladora del proyecto",
        "Análisis, diseño, implementación y pruebas",
        "Producto funcional documentado y verificado para titulación",
    ],
    [
        "Iglesia Evangélica La Alborada (IECA)",
        "Institución propietaria",
        "Control financiero por ministerio y aportación iglesia 33 % en talento",
    ],
]

TABLA_INDIRECTOS = [
    ["Beneficiario indirecto", "Relación con el proyecto", "Beneficio obtenido"],
    [
        "Feligresía y comunidad de IECA",
        "Zona de influencia institucional",
        "Mayor confianza en el uso de ofrendas y recursos ministeriales",
    ],
    [
        "Liderazgo pastoral",
        "Toma de decisiones institucionales",
        "Reportes consolidados y saldos por ministerio",
    ],
    [
        "Ministerios en general",
        "Destinatarios de asignaciones",
        "Visibilidad del saldo disponible y del kardex",
    ],
    [
        "Auditores o revisores internos",
        "Supervisión de fondos",
        "Historial exportable CSV/JSON y auditoría de movimientos",
    ],
    [
        "Futuros administradores y contables",
        "Sucesores en el cargo",
        "Sistema documentado con respaldos y manual de despliegue",
    ],
]

TABLA_MATRIZ = [
    ["Involucrado", "Interés", "Influencia", "Estrategia"],
    [
        "Administrador IECA",
        "Alto",
        "Alto",
        "Validación de requisitos y pruebas de aceptación",
    ],
    [
        "Contable IECA",
        "Alto",
        "Medio",
        "Validación de reportes y exportación Excel",
    ],
    [
        "Colaboradores",
        "Medio",
        "Bajo",
        "Pruebas de registro de movimientos pendientes",
    ],
    [
        "Desarrolladora",
        "Alto",
        "Alto",
        "Diseño, implementación y verificación",
    ],
    [
        "Comunidad / feligresía",
        "Medio",
        "Bajo",
        "Beneficio indirecto por transparencia financiera",
    ],
]


def set_run_font(run, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def find_paragraph_index(doc: Document, text: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    raise ValueError(f"Párrafo no encontrado: {text!r}")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def add_para_after_element(doc: Document, ref_element, text: str, bold: bool = False):
    p = doc.add_paragraph(style=STYLE_BODY)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    ref_element.addnext(p._element)
    return p._element


def add_caption_after(doc: Document, ref_element, text: str):
    return add_para_after_element(doc, ref_element, text, bold=True)


def add_table_after(doc: Document, ref_element, rows: list[list[str]]):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, bold=(r == 0), size=10)
    ref_element.addnext(table._element)
    return table._element


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup = DOC_PATH.with_suffix(
        f".backup-beneficiarios-{datetime.now():%Y%m%d-%H%M%S}.docx"
    )
    shutil.copy2(DOC_PATH, backup)
    print(f"Respaldo: {backup}")

    doc = Document(str(DOC_PATH))
    start = find_paragraph_index(doc, HEADING)
    end = find_paragraph_index(doc, END_HEADING)

    # Eliminar placeholders entre el título y Entregables (sin borrar los headings)
    to_remove = []
    for i in range(start + 1, end):
        to_remove.append(doc.paragraphs[i])
    for p in to_remove:
        remove_paragraph(p)

    ref = doc.paragraphs[start]._element

    ref = add_para_after_element(doc, ref, INTRO)
    ref = add_para_after_element(doc, ref, P_DIRECTOS)
    ref = add_caption_after(doc, ref, "Tabla 27. Beneficiarios directos del proyecto IECA")
    ref = add_table_after(doc, ref, TABLA_DIRECTOS)
    ref = add_para_after_element(doc, ref, P_INDIRECTOS)
    ref = add_caption_after(doc, ref, "Tabla 28. Beneficiarios indirectos del proyecto IECA")
    ref = add_table_after(doc, ref, TABLA_INDIRECTOS)
    ref = add_para_after_element(doc, ref, P_MARCO)
    ref = add_caption_after(doc, ref, "Tabla 29. Matriz de involucrados del proyecto IECA")
    ref = add_table_after(doc, ref, TABLA_MATRIZ)

    out = DOC_PATH
    try:
        doc.save(str(out))
        print(f"Guardado: {out}")
    except PermissionError:
        doc.save(str(OUT_FALLBACK))
        print(f"Archivo abierto en Word. Guardado en: {OUT_FALLBACK}")


if __name__ == "__main__":
    main()
