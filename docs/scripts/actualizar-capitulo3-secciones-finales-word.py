"""
Actualiza CAPITULO 3 - ENTENDERLO.docx con:
- Beneficiarios directos e indirectos
- Entregables del proyecto (EDT)
- Propuesta
"""
from __future__ import annotations

import importlib.util
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

SCRIPT_DIR = Path(__file__).resolve().parent


def load_module(name: str, filename: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPT_DIR / filename)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


_ben = load_module("beneficiarios_contenido", "beneficiarios-contenido.py")
_ent = load_module("entregables_propuesta_contenido", "entregables-propuesta-contenido.py")

DOC_PATH = Path(r"c:\Users\ena\Desktop\TITULACION\CAPITULO 3 - ENTENDERLO.docx")
OUT_FALLBACK = DOC_PATH.with_name(DOC_PATH.stem + "-actualizado.docx")

SECTION_HEADINGS = (
    _ben.HEADING,
    _ent.HEADING_ENTREGABLES,
    _ent.HEADING_PROPUESTA,
)
BASE_TABLE_COUNT = 12


def set_run_font(run, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(text), bold=bold, size=size)


def add_heading(doc: Document, text: str, level: int = 3) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, size=12)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), bold=bold, size=11)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val, bold=(r == 0), size=9)
    doc.add_paragraph()


def find_first_heading(doc: Document) -> int | None:
    for heading in SECTION_HEADINGS:
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip() == heading:
                return i
    return None


def remove_tail_sections(doc: Document) -> None:
    start = find_first_heading(doc)
    if start is None:
        return
    for i in range(len(doc.paragraphs) - 1, start - 1, -1):
        doc.paragraphs[i]._element.getparent().remove(doc.paragraphs[i]._element)
    while len(doc.tables) > BASE_TABLE_COUNT:
        doc.tables[-1]._tbl.getparent().remove(doc.tables[-1]._tbl)


def append_all_sections(doc: Document) -> None:
    add_heading(doc, _ben.HEADING, level=3)
    for text in _ben.PARAGRAPHS:
        add_para(doc, text)

    add_heading(doc, _ent.HEADING_ENTREGABLES, level=3)
    add_para(doc, _ent.ENTREGABLES_INTRO)
    add_para(doc, _ent.TABLA_EDT_CAPTION, bold=True)
    add_table(doc, _ent.TABLA_EDT)

    add_heading(doc, _ent.HEADING_PROPUESTA, level=3)
    for text in _ent.PROPUESTA_PARAGRAPHS:
        add_para(doc, text)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup = DOC_PATH.with_suffix(
        f".backup-secciones-{datetime.now():%Y%m%d-%H%M%S}.docx"
    )
    shutil.copy2(DOC_PATH, backup)
    print(f"Respaldo: {backup}")

    doc = Document(str(DOC_PATH))
    remove_tail_sections(doc)
    append_all_sections(doc)

    try:
        doc.save(str(DOC_PATH))
        print(f"Guardado: {DOC_PATH}")
    except PermissionError:
        doc.save(str(OUT_FALLBACK))
        print(f"Archivo abierto en Word. Guardado en: {OUT_FALLBACK}")
        print("Cierra Word y vuelve a ejecutar este script, o renombra el -actualizado.docx")


if __name__ == "__main__":
    main()
