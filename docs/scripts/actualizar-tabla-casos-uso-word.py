"""
Actualiza la tabla de casos de uso (4 módulos, versión simplificada)
en CAPITULO 3 - ENTENDERLO.docx
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

DOC_PATH = Path(r"c:\Users\ena\Desktop\TITULACION\CAPITULO 3 - ENTENDERLO.docx")
OUT_PATH = DOC_PATH  # si está abierto en Word, guarda en -actualizado.docx
TABLE_INDEX = 2

M1 = "Seguridad y Acceso (Autenticación)"
M2 = "Gestión Financiera y Flujo de Caja"
M3 = "Reportes y Analítica"
M4 = "Administración del Sistema y Soporte"

# ID, caso de uso, módulo, actor, descripción breve
CASOS_USO = [
    ("CU-01", "Iniciar sesión", M1, "Administrador, Contable, Colaborador", "Login con JWT y carga inicial de datos"),
    ("CU-02", "Recuperar contraseña", M1, "Administrador, Contable, Colaborador", "Código por correo y restablecimiento"),
    ("CU-03", "Cambiar contraseña", M1, "Administrador, Contable, Colaborador", "Cambio obligatorio en primer acceso"),
    ("CU-04", "Registrar ingreso", M2, "Colaborador, Administrador", "Alta con comprobante; queda pendiente si es colaborador"),
    ("CU-05", "Registrar gasto", M2, "Colaborador, Administrador", "Alta con comprobante; queda pendiente si es colaborador"),
    ("CU-06", "Aprobar movimiento", M2, "Administrador", "Aprueba ingreso o gasto; genera aportación 33% si es talento"),
    ("CU-07", "Rechazar movimiento", M2, "Administrador", "Rechaza con motivo"),
    ("CU-08", "Consultar dashboard", M3, "Administrador, Contable, Colaborador", "KPIs y gráficos con movimientos aprobados"),
    ("CU-09", "Generar reportes", M3, "Administrador, Contable, Colaborador", "Filtros, desglose, kardex y exportación Excel"),
    ("CU-10", "Gestionar ministerios", M4, "Administrador", "CRUD, saldo disponible y kardex"),
    ("CU-11", "Gestionar usuarios", M4, "Administrador", "CRUD con roles y asignación de ministerio"),
    ("CU-12", "Ejecutar cierre mensual", M4, "Administrador", "Cierra periodo y bloquea movimientos del mes"),
    ("CU-13", "Backup y restauración", M4, "Administrador", "Exportar e importar respaldo JSON"),
    ("CU-14", "Consultar auditoría", M4, "Administrador", "Historial de movimientos y exportación CSV"),
    ("CU-15", "Enviar alertas por correo", M4, "Administrador", "Pendientes antiguos y aviso de cierre"),
    ("CU-16", "Consultar notificaciones", M1, "Administrador, Contable, Colaborador", "Ver y marcar alertas en la interfaz"),
    ("CU-17", "Consultar movimientos generales", M2, "Contable", "Consulta de ingresos y gastos en modo solo lectura"),
    ("CU-18", "Consultar mis movimientos", M2, "Colaborador", "Consulta de movimientos de su ministerio"),
]

MODULOS_BULLETS = [
    "1. Seguridad y Acceso (Autenticación) — Actores: Administrador, Contable, Colaborador — CU-01, CU-02, CU-03, CU-16",
    "2. Gestión Financiera y Flujo de Caja — Actores: Colaborador, Administrador, Contable — CU-04, CU-05, CU-06, CU-07, CU-17, CU-18",
    "3. Reportes y Analítica — Actores: Administrador, Contable, Colaborador — CU-08, CU-09",
    "4. Administración del Sistema y Soporte — Actor: Administrador — CU-10 a CU-15",
]


def set_run_font(run, name: str = "Calibri", size_pt: float = 10, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size_pt)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_text(cell, text: str, bold: bool = False, size_pt: float = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, bold=bold, size_pt=size_pt)


def ensure_row_count(table, needed_rows: int) -> None:
    while len(table.rows) < needed_rows:
        table.add_row()
    while len(table.rows) > needed_rows:
        table._tbl.remove(table.rows[-1]._tr)


def update_module_bullets(doc: Document) -> None:
    org_para = None
    bullet_paras = []
    collecting = False
    for para in doc.paragraphs:
        if "Organización por módulos" in para.text:
            org_para = para
            collecting = True
            continue
        if collecting:
            t = para.text.strip()
            if t.startswith("Figura"):
                break
            if t and t[0].isdigit() and "CU-" in t:
                bullet_paras.append(para)

    if not org_para:
        return

    for p in bullet_paras:
        p._element.getparent().remove(p._element)

    ref = org_para._element
    for bullet in MODULOS_BULLETS:
        new_p = doc.add_paragraph(bullet, style="List Bullet")
        ref.addnext(new_p._element)
        ref = new_p._element


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = DOC_PATH.with_name(f"{DOC_PATH.stem}.backup-{stamp}.docx")
    shutil.copy2(DOC_PATH, backup)
    print(f"Backup: {backup}")

    doc = Document(str(DOC_PATH))
    table = doc.tables[TABLE_INDEX]

    headers = ["ID", "Caso de uso", "Módulo", "Actor principal", "Descripción breve"]
    needed = 1 + len(CASOS_USO)
    ensure_row_count(table, needed)

    for col, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[col], header, bold=True, size_pt=10)

    for row_idx, caso in enumerate(CASOS_USO, start=1):
        for col, value in enumerate(caso):
            set_cell_text(table.rows[row_idx].cells[col], value, size_pt=9)

    update_module_bullets(doc)

    try:
        doc.save(str(OUT_PATH))
        print(f"Actualizado: {OUT_PATH}")
    except PermissionError:
        alt = DOC_PATH.with_name(f"{DOC_PATH.stem}-actualizado.docx")
        doc.save(str(alt))
        print(f"Archivo original abierto. Guardado en: {alt}")
    print(f"Casos de uso: {len(CASOS_USO)} (4 módulos)")


if __name__ == "__main__":
    main()
