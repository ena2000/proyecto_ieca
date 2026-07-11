# -*- coding: utf-8 -*-
"""Genera .docx con Resultados, Criterios validacion/aceptacion y Conclusiones."""
from docx import Document
from docx.shared import Pt, RGBColor

OUT = r"C:\Users\ena\Desktop\TITULACION\CAP3-CAP4-PARA-COPIAR.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


def instruccion(texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def titulo(texto, nivel=2):
    doc.add_heading(texto, level=nivel)


def parrafo(texto):
    doc.add_paragraph(texto)


def viñeta(texto):
    doc.add_paragraph(texto, style="List Bullet")


# ===== PORTADA =====
titulo("Cap\u00edtulo III y IV \u2014 textos para copiar", 1)
instruccion(
    "Borra el texto de PLANTILLA en cada secci\u00f3n (incluido el ejemplo de los ratones, "
    "Delphi/Kendall y \u00abDetalle la conclusi\u00f3n 1\u00bb) y pega el contenido correspondiente. "
    "Mant\u00e9n el formato de vi\u00f1etas donde aplique."
)

# ===== 1) CRITERIOS DE VALIDACION =====
titulo("1) Criterios de validaci\u00f3n de la propuesta (Cap\u00edtulo III)")
instruccion(
    "Ubicaci\u00f3n: despu\u00e9s de \u00abPropuesta\u00bb, antes de \u00abResultados\u00bb. "
    "Borra desde \u00abDescriba el criterio y estrategia...\u00bb hasta antes del t\u00edtulo \u00abResultados\u00bb."
)

parrafo(
    "La validaci\u00f3n del prototipo no utiliz\u00f3 juicio de expertos ni el M\u00e9todo Delphi, "
    "sino verificaci\u00f3n t\u00e9cnica del software y pruebas de aceptaci\u00f3n con usuarios "
    "institucionales de IECA (administrador, contable y colaboradores de ministerio), en coherencia "
    "con los criterios \u00e9ticos del Anexo 5 y la estrategia documentada en el Anexo 7."
)
parrafo("La estrategia de validaci\u00f3n comprendi\u00f3 las siguientes acciones:")
viñeta(
    "Diagn\u00f3stico previo mediante entrevista semiestructurada y encuesta estructurada de siete "
    "\u00edtems (n = 5), cuyos resultados se presentan en las Tablas 14 a 21."
)
viñeta(
    "Pruebas unitarias del frontend: 46 casos con Karma y Jasmine (Tabla 30 y Tabla 31)."
)
viñeta(
    "Pruebas del backend: 49 casos con Node.js test runner y Supertest; 48 obligatorios en verde "
    "(Tabla 31 y Tabla 32)."
)
viñeta(
    "Pruebas de integraci\u00f3n HTTP sobre la API REST Express (Tabla 33)."
)
viñeta(
    "Pruebas de reglas de negocio: aportaci\u00f3n del 33 %, roles, cierre mensual y kardex (Tabla 34)."
)
viñeta(
    "Pruebas de aceptaci\u00f3n manual con ocho casos CP-M01 a CP-M08, ejecutados por administrador, "
    "contable y colaborador (Tabla 35)."
)
viñeta(
    "Integraci\u00f3n continua en GitHub Actions (lint, pruebas y build) en cada cambio del repositorio."
)
parrafo(
    "Criterio de validez: el prototipo se considera validado cuando las pruebas automatizadas "
    "cr\u00edticas superan en el entorno de integraci\u00f3n continua, los ocho casos manuales quedan "
    "verificados y la regla de aportaci\u00f3n del 33 % se comprueba en c\u00f3digo y en el caso real del "
    "ingreso de talento del ministerio Alabanza y Adoraci\u00f3n (Tablas 27 a 29)."
)

# ===== 2) RESULTADOS =====
titulo("2) Resultados (Cap\u00edtulo III)")
instruccion(
    "Ubicaci\u00f3n: secci\u00f3n \u00abResultados\u00bb, antes del Cap\u00edtulo IV. "
    "Borra todo el texto de plantilla (incluido el ejemplo de los ratones)."
)

parrafo("Resultados del diagn\u00f3stico")
parrafo(
    "La encuesta aplicada a la muestra intencional (n = 5: tres colaboradores de ministerio, "
    "un administrador y un contable) evidenci\u00f3 que el 100 % registra ingresos y gastos con "
    "m\u00e9todos manuales o dispersos: 80 % en hoja de c\u00e1lculo y 20 % en cuaderno; ning\u00fan "
    "informante dispone de un sistema digital integrado (Tabla 15, Figura 5)."
)
parrafo(
    "El 100 % conoce la regla de aportaci\u00f3n del 33 % a la iglesia (Tabla 16, Figura 6), "
    "por lo que la problem\u00e1tica operativa no se explica por desconocimiento normativo. "
    "En cuanto a la frecuencia de registro, el 60 % lo hace semanalmente, el 20 % solo en "
    "reuniones o cierres y el 20 % de forma irregular (Tabla 18, Figura 8); no existe registro "
    "en tiempo real."
)
parrafo(
    "El 20 % de los encuestados report\u00f3 p\u00e9rdida, duplicidad o desorden en la informaci\u00f3n "
    "financiera con el m\u00e9todo actual (Tabla 19, Figura 9). La totalidad calific\u00f3 con el "
    "m\u00e1ximo (5) la necesidad de una herramienta digital y el 100 % consider\u00f3 necesario un "
    "sistema web para centralizar ingresos y gastos por ministerio (Tablas 20 y 21, Figuras 10 y 11)."
)

parrafo("Resultados de la validaci\u00f3n t\u00e9cnica del prototipo")
parrafo(
    "El prototipo super\u00f3 94 de 95 pruebas automatizadas: 46/46 en frontend (Karma + Jasmine) "
    "y 48/49 en backend (Node test runner). El \u00fanico caso condicional corresponde al env\u00edo "
    "de alertas por SMTP, dependiente de configuraci\u00f3n local y sin impacto en la l\u00f3gica de "
    "negocio (Tablas 30 a 32)."
)
parrafo(
    "Las pruebas de integraci\u00f3n HTTP verificaron endpoints de autenticaci\u00f3n, ingresos, "
    "gastos y cierre mensual con resultado superado en todos los casos cr\u00edticos (Tabla 33). "
    "Las reglas de negocio \u2014aportaci\u00f3n del 33 % en cuenta 4105, aprobaci\u00f3n exclusiva del "
    "administrador, bloqueo por cierre mensual y coherencia del kardex\u2014 fueron verificadas "
    "(Tabla 34)."
)
parrafo(
    "Los ocho casos de aceptaci\u00f3n manual (CP-M01 a CP-M08) quedaron en estado verificado con "
    "administrador, contable y colaborador, cubriendo login, registro y aprobaci\u00f3n de movimientos, "
    "consulta de reportes, cierre mensual, exportaci\u00f3n Excel, kardex y carga inicial de datos "
    "(Tabla 35)."
)
parrafo(
    "En el caso de uso real del ingreso de talento por $182,00 en el ministerio Alabanza y Adoraci\u00f3n, "
    "el sistema gener\u00f3 autom\u00e1ticamente la aportaci\u00f3n del 33 % ($60,06) al ministerio General "
    "y registr\u00f3 el neto del 67 % ($121,94) en el kardex, con saldo final de $41,94 (Tablas 27 a 29)."
)

# ===== 3) CRITERIOS DE ACEPTACION =====
titulo("3) Criterios de aceptaci\u00f3n del producto o servicio (Cap\u00edtulo IV)")
instruccion(
    "Ubicaci\u00f3n: inicio del Cap\u00edtulo IV, secci\u00f3n \u00abCriterios de aceptaci\u00f3n del producto o servicio\u00bb. "
    "Borra el texto de plantilla gen\u00e9rico."
)

parrafo(
    "El prototipo de gesti\u00f3n financiera IECA se considera aceptado cuando cumple simult\u00e1neamente "
    "los criterios funcionales, no funcionales y de verificaci\u00f3n descritos a continuaci\u00f3n, "
    "evaluados con las pruebas documentadas en el Cap\u00edtulo III (Tablas 30 a 35) y la evidencia "
    "del caso real de aportaci\u00f3n del 33 % (Tablas 27 a 29)."
)

parrafo("Criterios funcionales de aceptaci\u00f3n:")
viñeta("CP-M01: Los tres roles (Administrador, Contable y Colaborador) inician sesi\u00f3n y acceden al dashboard seg\u00fan su perfil.")
viñeta("CP-M02: El colaborador registra un ingreso en estado pendiente con comprobante adjunto.")
viñeta("CP-M03: El administrador aprueba un ingreso de talento (cuenta 4105) y el sistema genera la aportaci\u00f3n del 33 % al ministerio General.")
viñeta("CP-M04: El contable consulta reportes y movimientos sin permisos de aprobaci\u00f3n ni edici\u00f3n.")
viñeta("CP-M05: El cierre mensual bloquea la edici\u00f3n de movimientos del periodo cerrado.")
viñeta("CP-M06: Administrador y contable exportan reportes a Excel con kardex por ministerio.")
viñeta("CP-M07: El kardex refleja saldos coherentes con los movimientos aprobados.")
viñeta("CP-M08: Tras el login, el sistema carga correctamente los datos iniciales (bootstrap) seg\u00fan el rol.")

parrafo("Criterios no funcionales de aceptaci\u00f3n:")
viñeta("Seguridad: autenticaci\u00f3n con JWT, contrase\u00f1as cifradas y acceso restringido por rol y ministerio.")
viñeta("Confiabilidad: 94 de 95 pruebas automatizadas superadas; reglas de negocio verificadas en Tabla 34.")
viñeta("Usabilidad: interfaz web orientada a escritorio, operable sin conocimientos t\u00e9cnicos avanzados.")
viñeta("Trazabilidad: auditor\u00eda de inicio de sesi\u00f3n, estados de movimiento y respaldo JSON del sistema.")

parrafo(
    "Resultado de la aceptaci\u00f3n: los ocho casos CP-M01 a CP-M08 fueron verificados con usuarios "
    "institucionales de IECA. El prototipo desplegado en Firebase Hosting y Render cumple los "
    "criterios definidos y queda disponible para operaci\u00f3n institucional."
)

# ===== 4) CONCLUSIONES =====
titulo("4) Conclusiones (Cap\u00edtulo IV)")
instruccion(
    "Ubicaci\u00f3n: secci\u00f3n \u00abConclusiones\u00bb. Borra el texto de plantilla "
    "(\u00abUna vez realizado el an\u00e1lisis...\u00bb, \u00abDetalle la conclusi\u00f3n 1\u00bb, etc.)."
)

INTRO_CONCLUSIONES = (
    "Una vez desarrollado e implementado el prototipo de aplicaci\u00f3n web para la gesti\u00f3n "
    "financiera de los ministerios de IECA, y validado mediante pruebas automatizadas, pruebas de "
    "aceptaci\u00f3n y el diagn\u00f3stico previo con usuarios institucionales, se enuncian las siguientes "
    "conclusiones en respuesta a los objetivos espec\u00edficos planteados:"
)
parrafo(INTRO_CONCLUSIONES)

CONCLUSIONES = [
    "Se dise\u00f1\u00f3 la arquitectura cliente-servidor del sistema y se definieron los requerimientos "
    "funcionales y no funcionales para el registro y control financiero por ministerio. La soluci\u00f3n "
    "qued\u00f3 especificada en 18 casos de uso, diagramas de arquitectura, clases y estados, y el "
    "modelo de datos Firestore (Tablas 22 y 23; Figuras 12 a 19; Anexo 9), considerando roles, "
    "flujo de aprobaci\u00f3n, reportes y la regla de aportaci\u00f3n del 33 %.",

    "Se implementaron los m\u00f3dulos de registro de ingresos y gastos con comprobantes, estados "
    "pendiente, aprobado y rechazado, y flujo de aprobaci\u00f3n por el administrador. El caso real del "
    "ingreso de talento por $182,00 demostr\u00f3 la generaci\u00f3n autom\u00e1tica de la aportaci\u00f3n "
    "institucional y el registro del monto neto en el ministerio correspondiente (Tablas 27 a 29).",

    "Se desarrollaron los m\u00f3dulos de consulta y visualizaci\u00f3n mediante dashboard, reportes con "
    "filtros, kardex por ministerio, cierre mensual y exportaci\u00f3n a Excel, facilitando el an\u00e1lisis "
    "de movimientos financieros. El saldo final del kardex del ministerio Alabanza y Adoraci\u00f3n "
    "qued\u00f3 en $41,94 tras la aprobaci\u00f3n del movimiento de talento.",

    "Se evalu\u00f3 el prototipo mediante 95 pruebas automatizadas (46 frontend y 49 backend) y ocho "
    "casos manuales de aceptaci\u00f3n (CP-M01 a CP-M08), verificando registro, consulta, aprobaci\u00f3n "
    "y generaci\u00f3n de reportes con administrador, contable y colaborador. El sistema super\u00f3 las "
    "pruebas cr\u00edticas y qued\u00f3 desplegado en producci\u00f3n (Tablas 30 a 35).",

    "El diagn\u00f3stico previo confirm\u00f3 la necesidad institucional de centralizar la informaci\u00f3n "
    "financiera: el 100 % de la muestra (n = 5) consider\u00f3 necesario un sistema web y calific\u00f3 "
    "con el m\u00e1ximo la utilidad de una herramienta digital (Tablas 20 y 21), lo cual fundamenta la "
    "pertinencia del prototipo desarrollado para IECA.",
]
for c in CONCLUSIONES:
    viñeta(c)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Recordatorio: ").bold = True
p.add_run(
    "Despu\u00e9s de pegar estos textos, actualiza tambi\u00e9n el Anexo 7 (cambiar \u00abValidaci\u00f3n de expertos\u00bb "
    "por \u00abValidaci\u00f3n del prototipo\u00bb) usando el contenido de docs/ANEXO-7-VALIDACION-PROTOTIPO.md."
)

doc.save(OUT)
print("[GUARDADO]", OUT)
