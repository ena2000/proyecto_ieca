# -*- coding: utf-8 -*-
"""Genera .docx con Recomendaciones y Trabajos futuros para Capitulo IV."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\ena\Desktop\TITULACION\RECOMENDACIONES-TRABAJOS-FUTUROS-PARA-COPIAR.docx"

INTRO_RECOMENDACIONES = (
    "Tras finalizar el desarrollo del prototipo de aplicaci\u00f3n web orientado a modernizar "
    "la gesti\u00f3n y el control de los registros financieros de los ministerios de la Iglesia "
    "del Evangelio Cuadrangular \u201cLa Alborada\u201d (IECA), empleando Ionic Framework, Angular, "
    "Node.js, Express, TypeScript y Firebase Firestore como tecnolog\u00edas de c\u00f3digo abierto, "
    "se presentan las siguientes recomendaciones estrat\u00e9gicas para optimizar el funcionamiento "
    "del sistema, fortalecer la transparencia financiera institucional y elevar la calidad de "
    "la adopci\u00f3n tecnol\u00f3gica en la operaci\u00f3n diaria de la congregaci\u00f3n:"
)

INTRO_TRABAJOS_FUTUROS = (
    "A partir de los resultados obtenidos, las limitaciones declaradas en el estudio y el "
    "potencial de ampliaci\u00f3n del prototipo desarrollado, se proponen las siguientes l\u00edneas "
    "de investigaci\u00f3n y desarrollo que podr\u00edan ejecutarse en etapas posteriores a la "
    "presente titulaci\u00f3n:"
)

RECOMENDACIONES = [
    "Implementar de forma permanente el prototipo en la operaci\u00f3n diaria de IECA, "
    "sustituyendo progresivamente el registro manual en Excel y cuadernos. La encuesta "
    "diagn\u00f3stica mostr\u00f3 que el 100 % de la muestra (n = 5) considera necesario un "
    "sistema web centralizado; por ello se recomienda que la administraci\u00f3n institucional "
    "adopte el panel como canal oficial de registro, consulta y aprobaci\u00f3n de movimientos.",

    "Capacitar a colaboradores, contable y administrador en el uso de los cuatro m\u00f3dulos "
    "del sistema (Seguridad y Acceso, Gesti\u00f3n Financiera, Reportes y Administraci\u00f3n), "
    "con \u00e9nfasis en el flujo pendiente-aprobado-rechazado, el kardex por ministerio y la "
    "regla de aportaci\u00f3n del 33 % en ingresos de talento (cuenta 4105). Esto reducir\u00e1 "
    "errores de digitaci\u00f3n y asegurar\u00e1 que los movimientos se registren con la frecuencia "
    "adecuada.",

    "Establecer una rutina institucional de cierre mensual, respaldo de datos y revisi\u00f3n "
    "de auditor\u00eda de inicio de sesi\u00f3n. El prototipo ya incorpora estas funciones; su uso "
    "peri\u00f3dico fortalecer\u00e1 la transparencia financiera y facilitar\u00e1 la rendici\u00f3n de cuentas "
    "ante el liderazgo pastoral.",

    "Mantener activa la integraci\u00f3n continua (GitHub Actions) y ejecutar pruebas "
    "automatizadas antes de cada actualizaci\u00f3n en producci\u00f3n. Las 95 pruebas documentadas "
    "en el Cap\u00edtulo III deben repetirse tras cambios en reglas de negocio, roles o "
    "estructura de Firestore, para preservar la confiabilidad del sistema.",

    "Ampliar la validaci\u00f3n con m\u00e1s usuarios institucionales una vez desplegado el sistema, "
    "superando la muestra intencional de la fase diagn\u00f3stica (n = 5). Se sugiere aplicar "
    "nuevamente los ocho casos de aceptaci\u00f3n manual (CP-M01 a CP-M08) con colaboradores "
    "de distintos ministerios y documentar observaciones de usabilidad para futuras mejoras.",

    "Configurar en producci\u00f3n el servicio SMTP para alertas por correo y notificaciones "
    "administrativas, ya que la prueba de env\u00edo de alertas depende de esa variable en el "
    "entorno local. Esto completar\u00e1 el m\u00f3dulo de comunicaci\u00f3n previsto en el dise\u00f1o del "
    "sistema.",
]

TRABAJOS_FUTUROS = [
    "Desarrollar la integraci\u00f3n del prototipo con sistemas contables externos o "
    "herramientas de facturaci\u00f3n electr\u00f3nica, aspecto expl\u00edcitamente excluido del alcance "
    "actual pero identificado como limitaci\u00f3n del estudio. Esta ampliaci\u00f3n permitir\u00eda "
    "conciliar autom\u00e1ticamente los movimientos registrados en IECA con la contabilidad "
    "formal de la instituci\u00f3n.",

    "Incorporar un m\u00f3dulo de planificaci\u00f3n y control presupuestario por ministerio, "
    "alineado con los antecedentes revisados sobre gesti\u00f3n presupuestaria (Tutiven Campos "
    "y Luna Rioja, 2023; Gracia y Camarena Rodr\u00edguez, 2025). El m\u00f3dulo podr\u00eda proyectar "
    "ingresos y egresos, comparar metas con la ejecuci\u00f3n real y alertar desviaciones antes "
    "del cierre mensual.",

    "Ampliar el esquema de roles y permisos m\u00e1s all\u00e1 de Administrador, Contable y "
    "Colaborador, habilitando permisos granulares por ministerio o por tipo de movimiento. "
    "Esto atender\u00eda escenarios de crecimiento institucional sin modificar la arquitectura "
    "cliente-servidor ya implementada.",

    "Realizar un estudio de usabilidad y satisfacci\u00f3n con una muestra mayor de la "
    "poblaci\u00f3n de IECA (N = 182), aplicando una encuesta post-implementaci\u00f3n que permita "
    "medir con mayor representatividad el impacto del sistema en la organizaci\u00f3n de la "
    "informaci\u00f3n financiera.",

    "Adaptar el prototipo como plantilla reutilizable para otras congregaciones o "
    "instituciones sin fines de lucro que requieran gesti\u00f3n financiera por departamentos, "
    "parametrizando ministerios, cuentas contables y reglas de aportaci\u00f3n seg\u00fan cada "
    "contexto organizacional.",
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

doc.add_heading("Recomendaciones y Trabajos futuros \u2014 Cap\u00edtulo IV", level=1)
p = doc.add_paragraph()
p.add_run(
    "Borra el texto de plantilla de la seccion Recomendaciones y Trabajos futuros "
    "y pega cada vi\u00f1eta. Mant\u00e9n el formato de lista con vi\u00f1etas del documento."
).italic = True

doc.add_heading("Recomendaciones", level=2)
doc.add_paragraph(INTRO_RECOMENDACIONES)
for texto in RECOMENDACIONES:
    doc.add_paragraph(texto, style="List Bullet")

doc.add_paragraph()
doc.add_heading("Trabajos futuros", level=2)
doc.add_paragraph(INTRO_TRABAJOS_FUTUROS)
for texto in TRABAJOS_FUTUROS:
    doc.add_paragraph(texto, style="List Bullet")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Nota: ").bold = True
p.add_run(
    "Estas recomendaciones y trabajos futuros se derivan del alcance, las limitaciones "
    "declaradas en el Capitulo I, los resultados de la encuesta (Tablas 15\u201321), las "
    "pruebas del prototipo (Tablas 30\u201335) y los beneficiarios directos e indirectos "
    "descritos en el Capitulo III."
)

doc.save(OUT)
print("[GUARDADO]", OUT)
print("Recomendaciones:", len(RECOMENDACIONES), "| Trabajos futuros:", len(TRABAJOS_FUTUROS))
