# -*- coding: utf-8 -*-
"""Genera un .docx con ABREVIATURAS y SIMBOLOGIA corregidas para la tesis IECA.

Solo incluye siglas/simbolos que REALMENTE aparecen en el documento.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

OUT = r"C:\Users\ena\Desktop\TITULACION\ABREVIATURAS-SIMBOLOGIA-PARA-COPIAR.docx"

# (sigla, significado)  -> orden alfabetico
ABREVIATURAS = [
    ("API", "Interfaz de Programaci\u00f3n de Aplicaciones (Application Programming Interface)"),
    ("CC.MM.FF", "Facultad de Ciencias Matem\u00e1ticas y F\u00edsicas"),
    ("C.I.", "C\u00e9dula de Identidad"),
    ("CI", "Integraci\u00f3n Continua (Continuous Integration)"),
    ("CRUD", "Crear, Leer, Actualizar y Eliminar (Create, Read, Update, Delete)"),
    ("CSS", "Hojas de Estilo en Cascada (Cascading Style Sheets)"),
    ("CSV", "Valores Separados por Comas (Comma-Separated Values)"),
    ("EDT", "Estructura de Desglose del Trabajo"),
    ("FCM", "Mensajer\u00eda en la Nube de Firebase (Firebase Cloud Messaging)"),
    ("HTML", "Lenguaje de Marcado de Hipertexto (HyperText Markup Language)"),
    ("HTTP", "Protocolo de Transferencia de Hipertexto (HyperText Transfer Protocol)"),
    ("HTTPS", "Protocolo Seguro de Transferencia de Hipertexto (HyperText Transfer Protocol Secure)"),
    ("IDE", "Entorno de Desarrollo Integrado (Integrated Development Environment)"),
    ("IECA", "Iglesia del Evangelio Cuadrangular \u201cLa Alborada\u201d"),
    ("Ing.", "Ingeniero"),
    ("ISO", "Organizaci\u00f3n Internacional de Normalizaci\u00f3n (International Organization for Standardization)"),
    ("JSON", "Notaci\u00f3n de Objetos de JavaScript (JavaScript Object Notation)"),
    ("JWT", "Token Web JSON (JSON Web Token)"),
    ("M. Sc.", "M\u00e1ster (Magister Scientiae)"),
    ("NoSQL", "Base de datos no relacional (Not only SQL)"),
    ("npm", "Gestor de Paquetes de Node (Node Package Manager)"),
    ("PDF", "Formato de Documento Port\u00e1til (Portable Document Format)"),
    ("REST", "Transferencia de Estado Representacional (Representational State Transfer)"),
    ("SDK", "Kit de Desarrollo de Software (Software Development Kit)"),
    ("SPA", "Aplicaci\u00f3n de P\u00e1gina \u00danica (Single-Page Application)"),
    ("UG", "Universidad de Guayaquil"),
    ("URL", "Localizador Uniforme de Recursos (Uniform Resource Locator)"),
]

# (simbolo, significado)
SIMBOLOGIA = [
    ("%", "Porcentaje"),
    ("$", "D\u00f3lar estadounidense (USD)"),
]

TAB_POS = Inches(1.6)  # posicion de la segunda columna (significado)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


def add_titulo(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.bold = True
    p.paragraph_format.space_after = Pt(12)


def add_fila(col1, col2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    # tab stop para alinear la segunda columna
    pf.tab_stops.add_tab_stop(TAB_POS, WD_TAB_ALIGNMENT.LEFT)
    p.add_run(col1 + "\t" + col2)


# ===== ABREVIATURAS =====
add_titulo("ABREVIATURAS")
for sigla, sig in ABREVIATURAS:
    add_fila(sigla, sig)

doc.add_paragraph()

# ===== SIMBOLOGIA =====
add_titulo("SIMBOLOG\u00cdA")
for sim, sig in SIMBOLOGIA:
    add_fila(sim, sig)

doc.save(OUT)
print("[GUARDADO]", OUT)
print("ABREVIATURAS:", len(ABREVIATURAS), "| SIMBOLOGIA:", len(SIMBOLOGIA))
