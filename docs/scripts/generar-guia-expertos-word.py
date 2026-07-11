# -*- coding: utf-8 -*-
"""Genera guia resumida para expertos — juicio de validacion del prototipo IECA."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\ena\Desktop\TITULACION\GUIA-JUICIO-EXPERTOS-IECA.docx"

TITULO_PROYECTO = (
    "Desarrollo de un prototipo de aplicaci\u00f3n web utilizando Ionic Framework "
    "para la gesti\u00f3n financiera de los ministerios de la Iglesia del Evangelio "
    "Cuadrangular \u201cLa Alborada\u201d (IECA)"
)

AUTORA = "Milena Nicole Mariscal Ponce"
CARRERA = "Software \u2014 Universidad de Guayaquil"
URL_SISTEMA = "https://gestion-ieca.web.app"
TIEMPO_ESTIMADO = "20 a 30 minutos"

CREDENCIALES = [
    ("Administrador", "milena.mariscal", "123456", "Aprueba movimientos, usuarios, cierre, backup"),
    ("Contable", "diznarda.quezada", "123456", "Consulta dashboard, ingresos, gastos y reportes"),
    ("Colaborador", "andres.quinde", "123456", "Registra ingresos/gastos de su ministerio (pendientes)"),
]

PASOS_DEMO = [
    "Abrir el sistema en navegador (Chrome o Edge, escritorio). Si tarda al cargar, esperar 30\u201360 s y recargar.",
    "Iniciar sesi\u00f3n como Administrador y revisar el Dashboard (KPIs y gr\u00e1ficos).",
    "Consultar Ingresos y Gastos: filtros, estados (pendiente/aprobado/rechazado) y comprobantes.",
    "Aprobar o rechazar un movimiento pendiente (solo Administrador).",
    "Ir a Reportes: filtrar por periodo/ministerio, revisar kardex y exportar Excel.",
    "Revisar Ministerios: saldo disponible y kardex del ministerio.",
    "Abrir Administraci\u00f3n: backup, auditor\u00eda CSV y alertas (solo visualizar; no cerrar mes real).",
    "Cerrar sesi\u00f3n e ingresar como Colaborador: registrar un ingreso o gasto en estado pendiente.",
    "Volver como Administrador y aprobar el movimiento registrado por el colaborador.",
    "Ingresar como Contable y verificar que solo puede consultar (no aprobar ni administrar).",
]

CRITERIOS = [
    ("Claridad", "Lenguaje, interfaz y comprensi\u00f3n del sistema"),
    ("Objetividad", "Funciones observables y medibles"),
    ("Actualidad", "Tecnolog\u00edas actuales (web, API REST, nube)"),
    ("Suficiencia", "Cantidad y calidad de funciones para IECA"),
    ("Intencionalidad", "Adecuaci\u00f3n al problema financiero por ministerio"),
    ("Consistencia", "Coherencia entre roles, flujos y reportes"),
    ("Metodolog\u00eda", "Relaci\u00f3n con el dise\u00f1o y requisitos del proyecto"),
    ("Aplicabilidad", "Facilidad de uso en la operaci\u00f3n institucional"),
]


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

# Margenes utiles para imprimir
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

title = doc.add_heading("Gu\u00eda para juicio de expertos", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph(TITULO_PROYECTO)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].bold = True

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f"Autora: {AUTORA}\n{CARRERA}").italic = True

doc.add_paragraph()

doc.add_heading("1. Objetivo de la validaci\u00f3n", level=1)
doc.add_paragraph(
    "Solicitamos su apoyo como experto(a) para validar el prototipo web de gesti\u00f3n "
    "financiera IECA. La evaluaci\u00f3n consiste en revisar el sistema en funcionamiento, "
    "completar el instrumento del Anexo 7 de la tesis (8 criterios con puntaje de 5 a 100) "
    f"y firmar el formulario. Tiempo estimado: {TIEMPO_ESTIMADO}."
)

doc.add_heading("2. Acceso al sistema", level=1)
p = doc.add_paragraph()
p.add_run("URL: ").bold = True
p.add_run(URL_SISTEMA)
doc.add_paragraph(
    "Requisitos: navegador actualizado (Chrome, Edge o Firefox), conexi\u00f3n a internet y "
    "uso recomendado en computadora de escritorio."
)
doc.add_paragraph(
    "Nota: si la primera carga demora, es normal (servidor en la nube). Espere y vuelva a intentar."
)

doc.add_heading("3. Credenciales de prueba", level=1)
add_table(doc, ["Rol", "Usuario", "Contrase\u00f1a", "Uso principal"], CREDENCIALES)

doc.add_heading("4. Qu\u00e9 evaluar (resumen funcional)", level=1)
doc.add_paragraph(
    "El prototipo centraliza ingresos y gastos por ministerio, con tres roles "
    "(Administrador, Contable, Colaborador) y cuatro m\u00f3dulos:"
)
for item in [
    "Seguridad y acceso: login, recuperar/cambiar contrase\u00f1a, notificaciones.",
    "Gesti\u00f3n financiera: registro de ingresos y gastos, comprobantes, aprobaci\u00f3n/rechazo.",
    "Reportes y anal\u00edtica: dashboard, filtros, kardex, exportaci\u00f3n a Excel.",
    "Administraci\u00f3n: ministerios, usuarios, cierre mensual, backup, auditor\u00eda.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph(
    "Reglas clave: solo los movimientos aprobados afectan saldos y reportes; los colaboradores "
    "registran en estado pendiente; en ingresos de talento (cuenta 4105) se aplica autom\u00e1ticamente "
    "el 33 % de aportaci\u00f3n a la iglesia."
)

doc.add_heading("5. Ruta sugerida de revisi\u00f3n", level=1)
for i, paso in enumerate(PASOS_DEMO, 1):
    doc.add_paragraph(f"{i}. {paso}", style="List Number")

doc.add_heading("6. Instrumento Anexo 7 \u2014 criterios a calificar", level=1)
doc.add_paragraph(
    "Marque un valor entre 5 y 100 para cada indicador (seg\u00fan el formulario entregado):"
)
add_table(doc, ["Indicador", "Qu\u00e9 valora"], CRITERIOS)

doc.add_heading("7. Datos que debe devolver el experto", level=1)
for item in [
    "Nombres y apellidos completos.",
    "T\u00edtulo profesional.",
    "C\u00e9dula de identidad (C.I.).",
    "Puntajes de los 8 criterios.",
    "Firma en el formulario del Anexo 7.",
    "Observaciones u opiniones (opcional, 2\u20133 l\u00edneas).",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("8. C\u00e1lculo del porcentaje de validaci\u00f3n", level=1)
doc.add_paragraph(
    "Por experto: sume los 8 puntajes, divida entre 800 y multiplique por 100. "
    "Ejemplo: si los puntajes suman 690 \u2192 690/800 = 86,25 %. "
    "Con varios expertos, se promedia el porcentaje de cada uno."
)

doc.add_paragraph()
cierre = doc.add_paragraph()
cierre.add_run("Contacto estudiante: ").bold = True
cierre.add_run("Milena Nicole Mariscal Ponce \u2014 Tel. 0956857445 \u2014 milena.mariscalp@ug.edu.ec")

doc.add_paragraph()
nota = doc.add_paragraph()
nota.add_run("Nota: ").bold = True
nota.add_run(
    "Este documento es una gu\u00eda operativa para la sesi\u00f3n de validaci\u00f3n. "
    "El instrumento oficial y la constancia firmada corresponden al Anexo 7 de la tesis."
)

doc.save(OUT)
print("[GUARDADO]", OUT)
