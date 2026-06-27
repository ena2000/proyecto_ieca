"""
Genera docs/RESULTADOS-PRUEBAS-PRODUCCION.docx con el resumen del checklist A→G.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parents[2]
OUT_PATH = ROOT / "docs" / "RESULTADOS-PRUEBAS-PRODUCCION.docx"


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(10)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx].cells[c_idx], val)
    doc.add_paragraph()


def add_section_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Resultados de pruebas en producción")
    tr.bold = True
    tr.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Sistema Web de Gestión Financiera — IECA")
    sr.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}\n"
        "Entorno: https://gestion-ieca.web.app\n"
        "API: https://ieca-api.onrender.com/api"
    )
    mr.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph(
        "Usuarios de prueba: Admin milena.mariscal · Contable diznarda.quezada · "
        "Colaborador andres.quinde (ministerio Adolescentes)."
    )

    add_section_title(doc, "1. Resumen ejecutivo")
    add_table(
        doc,
        ["Concepto", "Resultado"],
        [
            ["Checklist manual A → G", "57/57 aprobado — validación funcional completa"],
            ["Pruebas automatizadas (local)", "Frontend 46/46 · Backend 48/49 (1 fallo SMTP local esperado)"],
            ["Deploy frontend Firebase", "Realizado"],
            ["Inconveniente recurrente", "Render Free: cold start, demoras, mensaje «despertando»"],
            ["Pendiente infraestructura", "Render Starter + SMTP (botón web Enviar resumen) — cuando se pueda pagar"],
            ["Mejoras post-checklist", "P-01 y P-04 corregidas y desplegadas; P-02 y P-03 validadas por usuario"],
            ["Veredicto", "El sistema cumple los requisitos probados en producción con datos IECA"],
        ],
    )

    add_section_title(doc, "2. Sección A — Smoke test (10 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Rol", "Resultado", "Observaciones"],
        [
            ["A-01", "Web carga", "—", "Aprobado", "Login visible sin pantalla en blanco"],
            ["A-02", "Login incorrecto", "—", "Aprobado", "Mensaje correcto; usuario admin no existe"],
            ["A-03", "Login admin", "Admin", "Aprobado", "Dashboard y menú completo"],
            ["A-04", "Login contable", "Contable", "Aprobado", "Sin Usuarios, Ministerios ni Administración"],
            ["A-05", "Login colaborador", "Colaborador", "Aprobado", "Alcance solo su ministerio"],
            ["A-06", "Ingreso pendiente", "Colaborador", "Aprobado", "No suma al balance hasta aprobar"],
            ["A-07", "Aprobar ingreso", "Admin", "Aprobado", "Aparece en dashboard tras aprobar"],
            ["A-08", "Contable solo lectura", "Contable", "Aprobado", "Sin Aprobar/Rechazar"],
            ["A-09", "Exportar Excel", "Admin/Contable", "Aprobado", "Archivo .xlsx válido"],
            ["A-10", "Bootstrap tras login", "Cualquiera", "Aprobado", "GET /api/bootstrap OK (304 caché válido)"],
        ],
    )

    add_section_title(doc, "3. Sección B — Auth y sesión (5 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Resultado", "Observaciones"],
        [
            ["B-01", "Cerrar sesión", "Aprobado", "Vuelve a login"],
            ["B-02", "Usuario corto (< 4)", "Aprobado", "Validación en formulario"],
            ["B-03", "Contraseña corta (< 6)", "Aprobado", "Validación en formulario"],
            ["B-04", "Recuperar contraseña", "Aprobado", "Flujo OK; en prod depende de SMTP"],
            ["B-05", "Sesión expirada", "Aprobado", "Tras limpiar storage, sin errores en cascada"],
        ],
    )

    add_section_title(doc, "4. Sección C — Movimientos CRUD (12 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Rol", "Resultado", "Observaciones"],
        [
            ["C-01", "Crear ingreso", "Colaborador", "Aprobado", "Queda pendiente"],
            ["C-02", "Crear gasto", "Colaborador", "Aprobado", "Queda pendiente"],
            ["C-03", "Aprobar gasto", "Admin", "Aprobado", "Reflejado en reportes"],
            ["C-04", "Rechazar ingreso", "Admin", "Aprobado", "No entra al balance"],
            ["C-05", "Colaborador no aprueba", "Colaborador", "Aprobado", "Sin botones de aprobación"],
            ["C-06", "Contable no aprueba", "Contable", "Aprobado", "Solo consulta"],
            ["C-07", "Solo su ministerio", "Colaborador", "Aprobado", "Solo Adolescentes"],
            ["C-08", "Comprobante (foto)", "Colaborador", "Aprobado", "Imagen visible al abrir"],
            ["C-09", "Editar pendiente", "Colaborador", "Aprobado", "Sigue pendiente"],
            ["C-10", "Editar aprobado", "Admin", "Aprobado", "Actualiza kardex/reportes (mes abierto)"],
            ["C-11", "Eliminar pendiente", "Colaborador", "Aprobado", "Desaparece de lista"],
            ["C-12", "Eliminar aprobado", "Admin", "Aprobado", "Balance coherente"],
        ],
    )
    doc.add_paragraph(
        "Extra: orden por fecha (más recientes primero) en ingresos y gastos — corregido y validado."
    )

    add_section_title(doc, "5. Sección D — Aportación 33 % talento (5 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Resultado", "Observaciones / inconvenientes"],
        [
            ["D-01", "Ingreso talento 4105", "Aprobado", "Pendiente hasta aprobar"],
            ["D-02", "Al aprobar ($100)", "Aprobado", "$33 General · $67 neto ministerio"],
            ["D-03", "Caso demo Andrés $182", "Aprobado", "Kardex neto coherente (~$41,94)"],
            ["D-04", "33 % no editable", "Aprobado", "Movimiento automático bloqueado"],
            ["D-05", "Borrar origen", "Aprobado", "Elimina también aportación en General (cascada)"],
        ],
    )
    doc.add_paragraph(
        "Inconveniente P-01: el 33 % tardaba en listas/dashboard al aprobar. "
        "Corregido (vista optimista + merge) y desplegado en Firebase."
    )

    add_section_title(doc, "6. Sección E — Reportes y dashboard (9 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Resultado", "Observaciones"],
        [
            ["E-01", "KPIs solo aprobados", "Aprobado", ""],
            ["E-02", "Gráficos", "Aprobado", "Sin errores de carga"],
            ["E-03", "Filtro período", "Aprobado", "Totales cambian"],
            ["E-04", "Filtro ministerio", "Aprobado", "Desglose por ministerio corregido en sesión"],
            ["E-05", "Saldo disponible", "Aprobado", "Histórico correcto"],
            ["E-06", "Kardex", "Aprobado", "Saldo = ingresos − gastos aprobados"],
            ["E-07", "Columnas aportación", "Aprobado", "Período y acumulado visibles"],
            ["E-08", "Reportes colaborador", "Aprobado", "Solo su ministerio"],
            ["E-09", "Dashboard colaborador", "Aprobado", "Datos de su ministerio"],
        ],
    )
    doc.add_paragraph(
        "P-03 (reportes $0 al cargar): reportado al inicio; indicado como resuelto por el usuario."
    )

    add_section_title(doc, "7. Sección F — Administración / catálogo (13 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Resultado", "Observaciones / inconvenientes"],
        [
            ["F-01", "Ministerios demo", "Aprobado", "22 + General"],
            ["F-02", "Usuarios demo", "Aprobado", "~182 usuarios"],
            ["F-03", "Staff admin", "Aprobado", "Milena y Orbe"],
            ["F-04", "Staff contable", "Aprobado", "Diznarda"],
            ["F-05", "Crear ministerio", "Aprobado", "Alta OK"],
            ["F-06", "Duplicado ministerio", "Aprobado", "Error unicidad"],
            ["F-07", "Crear colaborador", "Aprobado", "Login del nuevo usuario OK"],
            ["F-08", "Email duplicado", "Aprobado", "Error unicidad"],
            ["F-09", "Backup JSON", "Aprobado", "Exportación OK"],
            ["F-10", "Auditoría + CSV", "Aprobado", "Export CSV OK"],
            [
                "F-11",
                "Alertas email",
                "Aprobado",
                "GitHub Actions + correo real OK; botón web en Render Free fallaba (SMTP) — plan Starter",
            ],
            ["F-12", "Resumen aportación 33 %", "Aprobado", "Totales por ministerio; filtro corregido"],
            ["F-13", "Contable sin /administracion", "Aprobado", "Redirección + toast permiso"],
        ],
    )
    doc.add_paragraph("P-04: parpadeo «ministerio ya existe» — corregido en código y deploy.")

    add_section_title(doc, "8. Sección G — Cierre mensual (3 pruebas)")
    add_table(
        doc,
        ["ID", "Prueba", "Resultado", "Observaciones"],
        [
            ["G-01", "Cerrar mes", "Aprobado", "Periodo marcado cerrado; botón deshabilitado"],
            ["G-02", "Bloqueo crear", "Aprobado", "Junio bloqueado; julio permitido"],
            ["G-03", "Bloqueo editar", "Aprobado", "Admin no edita/aprueba/elimina en mes cerrado"],
        ],
    )
    doc.add_paragraph(
        "Tras validar G, se restauró backup (~2 min) para reabrir junio. "
        "Al inicio apareció aviso Render despertando — normal en plan Free."
    )

    add_section_title(doc, "9. Mejoras post-checklist (P)")
    add_table(
        doc,
        ["ID", "Tema", "Resultado", "Acción"],
        [
            ["P-01", "Demora 33 % al aprobar", "Corregido + deploy", "Vista optimista + sync lista"],
            ["P-02", "Ocultar General en formularios", "Validado por usuario", "Sin cambios pendientes"],
            ["P-03", "Reportes $0 / reset ministerio", "Validado por usuario", "Sin cambios pendientes"],
            ["P-04", "Parpadeo ministerio duplicado", "Corregido + deploy", ""],
            ["—", "Filtros estado por ministerio", "Corregido", "Contadores por ministerio"],
            ["—", "Actividad: detalle vs ingreso #N", "Corregido", "Descripción con detalle del origen"],
            ["—", "Test encoding aportación", "Corregido", "Escapes Unicode en util"],
        ],
    )

    add_section_title(doc, "10. Infraestructura y automatización")
    add_table(
        doc,
        ["Tema", "Resultado", "Notas"],
        [
            ["Firebase Hosting (frontend)", "Desplegado", "Incluye P-01, P-04 y fixes recientes"],
            ["Render API (Free)", "Operativo con limitaciones", "Cold start 30–60 s; restore backup lento"],
            ["SMTP / alertas desde panel web", "Pendiente", "Requiere Render Starter + SMTP"],
            ["Alertas vía GitHub Actions", "Aprobado", "Resumen operativo por correo (F-11)"],
            ["Brevo API", "Descartado", "Eliminado del código; solo SMTP"],
        ],
    )

    add_section_title(doc, "11. Conteo final")
    add_table(
        doc,
        ["Sección", "Pruebas", "Aprobadas", "Omitidas", "Fallidas"],
        [
            ["A Smoke", "10", "10", "0", "0"],
            ["B Auth", "5", "5", "0", "0"],
            ["C Movimientos", "12", "12", "0", "0"],
            ["D Aportación", "5", "5", "0", "0"],
            ["E Reportes", "9", "9", "0", "0"],
            ["F Admin", "13", "13", "0", "0"],
            ["G Cierre", "3", "3", "0", "0"],
            ["TOTAL checklist", "57", "57", "0", "0"],
        ],
    )

    add_section_title(doc, "12. Inconvenientes globales (informe / tesis)")
    for item in [
        "Render Free: latencia inicial, mensajes «API despertando», restore backup lento (~2 min).",
        "SMTP en Free: envío desde botón web de Administración no fiable; workaround con GitHub Actions.",
        "Demora visual del 33 % (P-01): resuelto antes del cierre del proyecto.",
        "Cierre mensual irreversible: probado con backup previo y restauración exitosa.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    print(f"OK: {OUT_PATH}")


if __name__ == "__main__":
    main()
